"""
Backend/resume_render.py

One source of truth for how a built resume looks.

    build_document(data)  ->  ordered, cleaned structure
    render_text(data)     ->  plain text   (AI tailoring / ATS / cover letters)
    render_docx(data)     ->  .docx bytes  (download, with clickable links)
    render_html(data)     ->  HTML string  (live preview)

Because all three renderers read the same structure, the preview, the
download and the text the AI sees can never disagree about content or order.
"""

import html
import io
import re

# ------------------------------------------------------------------
# Sections + ordering
# ------------------------------------------------------------------

SECTION_TITLES = {
    "summary": "Summary",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
    "skills": "Skills",
    "certifications": "Certifications",
}
DEFAULT_ORDER = list(SECTION_TITLES)  # summary, experience, projects, ...


def resolve_order(data: dict) -> list[str]:
    """
    Turn whatever the client sent into a safe, complete section order:
    unknown keys dropped, duplicates dropped, missing sections appended.
    """
    order, seen = [], set()
    for key in data.get("section_order") or []:
        if key in SECTION_TITLES and key not in seen:
            order.append(key)
            seen.add(key)
    order += [k for k in DEFAULT_ORDER if k not in seen]
    return order


# ------------------------------------------------------------------
# Links
# ------------------------------------------------------------------

def normalize_url(raw) -> str | None:
    """
    'github.com/me'        -> 'https://github.com/me'
    'me@mail.com'          -> 'mailto:me@mail.com'
    'javascript:alert(1)'  -> None   (only http/https/mailto are allowed)
    """
    url = (raw or "").strip()
    if not url:
        return None
    if re.match(r"^(https?://|mailto:)", url, re.I):
        return url
    if re.fullmatch(r"[^@\s/]+@[^@\s/]+\.[^@\s/]+", url):
        return "mailto:" + url
    if re.match(r"^[a-z][a-z0-9+.\-]*:", url, re.I) and not re.match(r"^[\w.\-]+:\d+", url):
        return None  # some other scheme (javascript:, data:, file: ...)
    return "https://" + url


def _clean(s) -> str:
    return (s or "").strip()


def _link(text: str, url_raw: str | None, plain: str | None = None) -> dict:
    return {"text": text, "url": normalize_url(url_raw), "plain": plain or text}


# ------------------------------------------------------------------
# Structure
# ------------------------------------------------------------------

def _contact_items(data: dict) -> list[dict]:
    items = []
    if _clean(data.get("email")):
        e = _clean(data["email"])
        items.append(_link(e, e))
    for key in ("phone", "location"):
        if _clean(data.get(key)):
            items.append({"text": _clean(data[key]), "url": None, "plain": _clean(data[key])})
    for key in ("linkedin", "github", "portfolio"):
        v = _clean(data.get(key))
        if v:
            items.append(_link(v, v))
    for extra in data.get("links") or []:
        url = _clean(extra.get("url"))
        if not url:
            continue
        label = _clean(extra.get("label")) or re.sub(r"^https?://", "", url)
        items.append(_link(label, url, plain=f"{label}: {url}"))
    return items


def _entry(head_parts, meta="", link=None, detail="", bullets=None) -> dict:
    return {
        "head": [p for p in head_parts if _clean(p)],
        "meta": _clean(meta),
        "link": link,
        "detail": _clean(detail),
        "bullets": [_clean(b) for b in (bullets or []) if _clean(b)],
    }


def build_document(data: dict) -> dict:
    sections = {}

    if _clean(data.get("summary")):
        sections["summary"] = [{"kind": "text", "text": _clean(data["summary"])}]

    exp = [
        _entry([e.get("role"), e.get("company")], e.get("dates"), bullets=e.get("bullets"))
        for e in data.get("experience") or []
        if _clean(e.get("role")) or _clean(e.get("company"))
    ]
    if exp:
        sections["experience"] = [{"kind": "entry", **x} for x in exp]

    proj = []
    for p in data.get("projects") or []:
        if not _clean(p.get("title")):
            continue
        url = _clean(p.get("link"))
        link = _link("Link", url, plain=url) if url else None
        proj.append(_entry([p.get("title")], p.get("tech"), link=link, bullets=p.get("bullets")))
    if proj:
        sections["projects"] = [{"kind": "entry", **x} for x in proj]

    edu = [
        _entry([e.get("degree"), e.get("institution")], e.get("dates"), detail=e.get("details"))
        for e in data.get("education") or []
        if _clean(e.get("degree")) or _clean(e.get("institution"))
    ]
    if edu:
        sections["education"] = [{"kind": "entry", **x} for x in edu]

    skills = [_clean(s) for s in data.get("skills") or [] if _clean(s)]
    if skills:
        sections["skills"] = [{"kind": "text", "text": ", ".join(skills)}]

    certs = [_clean(c) for c in data.get("certifications") or [] if _clean(c)]
    if certs:
        sections["certifications"] = [{"kind": "bullet", "text": c} for c in certs]

    return {
        "name": _clean(data.get("full_name")),
        "contact": _contact_items(data),
        "sections": [
            {"key": k, "title": SECTION_TITLES[k], "items": sections[k]}
            for k in resolve_order(data)
            if k in sections
        ],
    }


# ------------------------------------------------------------------
# Plain text
# ------------------------------------------------------------------

def _entry_head_text(item: dict) -> str:
    head = " - ".join(item["head"])
    if item["meta"]:
        head += f" ({item['meta']})"
    if item["link"]:
        head += f" - {item['link']['plain']}"
    return head


def render_text(data: dict) -> str:
    doc = build_document(data)
    lines = [doc["name"]]
    if doc["contact"]:
        lines.append(" | ".join(c["plain"] for c in doc["contact"]))

    for sec in doc["sections"]:
        lines += ["", sec["title"].upper()]
        for item in sec["items"]:
            if item["kind"] == "text":
                lines.append(item["text"])
            elif item["kind"] == "bullet":
                lines.append(f"- {item['text']}")
            else:
                lines.append(_entry_head_text(item))
                if item["detail"]:
                    lines.append(item["detail"])
                lines += [f"- {b}" for b in item["bullets"]]
    return "\n".join(lines).strip()


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------

def _add_hyperlink(paragraph, text: str, url: str):
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_link_or_text(paragraph, link: dict):
    if link["url"]:
        _add_hyperlink(paragraph, link["text"], link["url"])
    else:
        paragraph.add_run(link["text"])


def render_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc_data = build_document(data)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(doc_data["name"])
    name_run.bold = True
    name_run.font.size = Pt(18)

    if doc_data["contact"]:
        p = doc.add_paragraph()
        for i, c in enumerate(doc_data["contact"]):
            if i:
                p.add_run(" | ")
            _add_link_or_text(p, c)

    for sec in doc_data["sections"]:
        doc.add_heading(sec["title"], level=2)
        for item in sec["items"]:
            if item["kind"] == "text":
                doc.add_paragraph(item["text"])
            elif item["kind"] == "bullet":
                doc.add_paragraph(item["text"], style="List Bullet")
            else:
                p = doc.add_paragraph()
                p.add_run(" — ".join(item["head"])).bold = True
                if item["meta"]:
                    p.add_run(f"  ({item['meta']})")
                if item["link"]:
                    p.add_run("  ")
                    _add_link_or_text(p, item["link"])
                if item["detail"]:
                    doc.add_paragraph(item["detail"])
                for b in item["bullets"]:
                    doc.add_paragraph(b, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# HTML preview
# ------------------------------------------------------------------

_PREVIEW_CSS = """
.resume-preview{font-family:Calibri,Carlito,Arial,sans-serif;font-size:10.5pt;line-height:1.4;color:#111;background:#fff;padding:32px 40px;max-width:780px;margin:0 auto}
.resume-preview h1{font-size:18pt;margin:0 0 4px}
.resume-preview h2{font-size:12.5pt;margin:16px 0 6px;padding-bottom:2px;border-bottom:1px solid #bbb;color:#1f3864}
.resume-preview p{margin:2px 0}
.resume-preview ul{margin:2px 0 6px 20px;padding:0}
.resume-preview a{color:#0563c1;text-decoration:underline}
.resume-preview .contact{margin-bottom:4px}
.resume-preview .sep{color:#888;margin:0 6px}
.resume-preview .entry{margin-top:6px}
.resume-preview .meta{color:#333}
""".strip()


def _h(s: str) -> str:
    return html.escape(s, quote=True)


def _link_html(link: dict) -> str:
    if link["url"]:
        return (
            f'<a href="{_h(link["url"])}" target="_blank" '
            f'rel="noopener noreferrer">{_h(link["text"])}</a>'
        )
    return _h(link["text"])


def render_html(data: dict) -> str:
    d = build_document(data)
    out = [f"<style>{_PREVIEW_CSS}</style>", '<div class="resume-preview">']
    out.append(f"<h1>{_h(d['name'])}</h1>")
    if d["contact"]:
        sep = '<span class="sep">|</span>'
        out.append(f'<p class="contact">{sep.join(_link_html(c) for c in d["contact"])}</p>')

    for sec in d["sections"]:
        out.append(f"<h2>{_h(sec['title'])}</h2>")
        bullets_only = [i for i in sec["items"] if i["kind"] == "bullet"]
        if bullets_only:
            out.append("<ul>" + "".join(f"<li>{_h(i['text'])}</li>" for i in bullets_only) + "</ul>")
            continue
        for item in sec["items"]:
            if item["kind"] == "text":
                out.append(f"<p>{_h(item['text'])}</p>")
                continue
            head = f"<strong>{_h(' — '.join(item['head']))}</strong>"
            if item["meta"]:
                head += f' <span class="meta">({_h(item["meta"])})</span>'
            if item["link"]:
                head += " " + _link_html(item["link"])
            out.append(f'<div class="entry"><p>{head}</p>')
            if item["detail"]:
                out.append(f"<p>{_h(item['detail'])}</p>")
            if item["bullets"]:
                out.append("<ul>" + "".join(f"<li>{_h(b)}</li>" for b in item["bullets"]) + "</ul>")
            out.append("</div>")

    out.append("</div>")
    return "\n".join(out)